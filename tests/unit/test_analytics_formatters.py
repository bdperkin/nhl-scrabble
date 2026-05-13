"""Unit tests for analytics formatters."""

import json
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path

import pytest

from nhl_scrabble.analytics.analyzer import CoverageGap, TestPerformance
from nhl_scrabble.analytics.formatters import (
    CSVFormatter,
    ExcelFormatter,
    HTMLFormatter,
    JSONFormatter,
    MarkdownFormatter,
    TableFormatter,
    TemplateFormatter,
    TextFormatter,
    XMLFormatter,
    YAMLFormatter,
)


class TestTextFormatter:
    """Tests for TextFormatter class."""

    @pytest.fixture
    def formatter(self) -> TextFormatter:
        """Create TextFormatter instance."""
        return TextFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/high_priority.py",
                current_coverage=50.0,
                target_coverage=90.0,
                lines_needed=40,
                priority="high",
            ),
            CoverageGap(
                module="src/medium_priority.py",
                current_coverage=75.0,
                target_coverage=90.0,
                lines_needed=15,
                priority="medium",
            ),
            CoverageGap(
                module="src/low_priority.py",
                current_coverage=85.0,
                target_coverage=90.0,
                lines_needed=5,
                priority="low",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_slow_operation",
                avg_duration=5.5,
                max_duration=10.0,
                min_duration=3.0,
                failure_rate=0.1,
                flakiness_score=0.8,
            ),
            TestPerformance(
                test_name="test_flaky_test",
                avg_duration=1.0,
                max_duration=2.0,
                min_duration=0.5,
                failure_rate=0.3,
                flakiness_score=0.95,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: TextFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test formatting coverage gaps."""
        result = formatter._format_coverage_gaps(sample_gaps)

        assert "Coverage Gaps" in result
        assert "src/high_priority.py" in result
        assert "50.0%" in result
        assert "90.0%" in result
        assert "HIGH" in result
        assert "MEDIUM" in result
        assert "LOW" in result

    def test_format_coverage_gaps_truncates_at_20(
        self,
        formatter: TextFormatter,
    ) -> None:
        """Test coverage gaps are truncated to 20 items."""
        gaps = [
            CoverageGap(
                module=f"module_{i}.py",
                current_coverage=50.0,
                target_coverage=90.0,
                lines_needed=40,
                priority="high",
            )
            for i in range(30)
        ]

        result = formatter._format_coverage_gaps(gaps)

        # Should only show first 20
        assert "module_0.py" in result
        assert "module_19.py" in result
        assert "module_20.py" not in result

    def test_format_slow_tests(
        self,
        formatter: TextFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test formatting slow tests."""
        result = formatter._format_slow_tests(sample_tests)

        assert "Slowest Tests" in result
        assert "test_slow_operation" in result
        assert "5.50s" in result
        assert "10.00s" in result
        assert "10.0%" in result  # failure rate

    def test_format_flaky_tests(
        self,
        formatter: TextFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test formatting flaky tests."""
        result = formatter._format_flaky_tests(sample_tests)

        assert "Flaky Tests" in result
        assert "test_flaky_test" in result
        assert "0.950" in result  # flakiness score
        assert "30.0%" in result  # failure rate

    def test_format_trends_improving(
        self,
        formatter: TextFormatter,
    ) -> None:
        """Test formatting improving coverage trend."""
        history = [
            {"coverage": 85.0, "commit": "abc123"},
            {"coverage": 82.0, "commit": "def456"},
            {"coverage": 80.0, "commit": "ghi789"},
        ]

        result = formatter._format_trends("improving", history)

        assert "Coverage Trends" in result
        assert "IMPROVING" in result
        assert "85.0%" in result
        assert "+5.0%" in result  # change from 80 to 85

    def test_format_trends_declining(
        self,
        formatter: TextFormatter,
    ) -> None:
        """Test formatting declining coverage trend."""
        history = [
            {"coverage": 80.0, "commit": "abc123"},
            {"coverage": 85.0, "commit": "def456"},
        ]

        result = formatter._format_trends("declining", history)

        assert "DECLINING" in result
        assert "80.0%" in result
        assert "-5.0%" in result

    def test_format_trends_stable(
        self,
        formatter: TextFormatter,
    ) -> None:
        """Test formatting stable coverage trend."""
        history = [
            {"coverage": 85.0, "commit": "abc123"},
            {"coverage": 85.5, "commit": "def456"},
        ]

        result = formatter._format_trends("stable", history)

        assert "STABLE" in result
        assert "85.0%" in result

    def test_format_trends_insufficient_data(
        self,
        formatter: TextFormatter,
    ) -> None:
        """Test formatting with insufficient data."""
        history = [{"coverage": 85.0, "commit": "abc123"}]

        result = formatter._format_trends("insufficient_data", history)

        assert "INSUFFICIENT_DATA" in result

    def test_format_trends_empty_history(
        self,
        formatter: TextFormatter,
    ) -> None:
        """Test formatting with empty history."""
        result = formatter._format_trends("insufficient_data", [])

        assert "0.0%" in result

    def test_format_complete_report(
        self,
        formatter: TextFormatter,
        sample_gaps: list[CoverageGap],
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test formatting complete analytics report."""
        data = {
            "coverage_gaps": sample_gaps,
            "slow_tests": sample_tests,
            "flaky_tests": sample_tests,
            "coverage_trend": "improving",
            "coverage_history": [
                {"coverage": 85.0, "commit": "abc"},
                {"coverage": 80.0, "commit": "def"},
            ],
        }

        result = formatter.format(data)

        # All sections should be present
        assert "Coverage Gaps" in result
        assert "Slowest Tests" in result
        assert "Flaky Tests" in result
        assert "Coverage Trends" in result

    def test_format_partial_data(
        self,
        formatter: TextFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test formatting with only some data fields."""
        data = {"coverage_gaps": sample_gaps}

        result = formatter.format(data)

        assert "Coverage Gaps" in result
        assert "Slowest Tests" not in result


class TestJSONFormatter:
    """Tests for JSONFormatter class."""

    @pytest.fixture
    def formatter(self) -> JSONFormatter:
        """Create JSONFormatter instance."""
        return JSONFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/test.py",
                current_coverage=75.0,
                target_coverage=90.0,
                lines_needed=15,
                priority="medium",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_example",
                avg_duration=2.5,
                max_duration=5.0,
                min_duration=1.0,
                failure_rate=0.05,
                flakiness_score=0.3,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: JSONFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test JSON formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        parsed = json.loads(result)
        assert "coverage_gaps" in parsed
        assert len(parsed["coverage_gaps"]) == 1
        assert parsed["coverage_gaps"][0]["module"] == "src/test.py"
        assert parsed["coverage_gaps"][0]["current_coverage"] == 75.0
        assert parsed["coverage_gaps"][0]["priority"] == "medium"

    def test_format_slow_tests(
        self,
        formatter: JSONFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test JSON formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        parsed = json.loads(result)
        assert "slow_tests" in parsed
        assert len(parsed["slow_tests"]) == 1
        assert parsed["slow_tests"][0]["test_name"] == "test_example"
        assert parsed["slow_tests"][0]["avg_duration"] == 2.5
        assert parsed["slow_tests"][0]["max_duration"] == 5.0
        assert parsed["slow_tests"][0]["failure_rate"] == 0.05

    def test_format_flaky_tests(
        self,
        formatter: JSONFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test JSON formatting of flaky tests."""
        data = {"flaky_tests": sample_tests}
        result = formatter.format(data)

        parsed = json.loads(result)
        assert "flaky_tests" in parsed
        assert len(parsed["flaky_tests"]) == 1
        assert parsed["flaky_tests"][0]["test_name"] == "test_example"
        assert parsed["flaky_tests"][0]["flakiness_score"] == 0.3

    def test_format_coverage_trend(
        self,
        formatter: JSONFormatter,
    ) -> None:
        """Test JSON formatting of coverage trend."""
        data = {"coverage_trend": "improving"}
        result = formatter.format(data)

        parsed = json.loads(result)
        assert "coverage_trend" in parsed
        assert parsed["coverage_trend"] == "improving"

    def test_format_coverage_history(
        self,
        formatter: JSONFormatter,
    ) -> None:
        """Test JSON formatting of coverage history."""
        history = [
            {"coverage": 85.0, "commit": "abc123"},
            {"coverage": 80.0, "commit": "def456"},
        ]
        data = {"coverage_history": history}
        result = formatter.format(data)

        parsed = json.loads(result)
        assert "coverage_history" in parsed
        assert len(parsed["coverage_history"]) == 2
        assert parsed["coverage_history"][0]["coverage"] == 85.0

    def test_format_complete_data(
        self,
        formatter: JSONFormatter,
        sample_gaps: list[CoverageGap],
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test JSON formatting of complete analytics data."""
        data = {
            "coverage_gaps": sample_gaps,
            "slow_tests": sample_tests,
            "flaky_tests": sample_tests,
            "coverage_trend": "stable",
            "coverage_history": [{"coverage": 85.0, "commit": "abc"}],
        }
        result = formatter.format(data)

        parsed = json.loads(result)
        assert "coverage_gaps" in parsed
        assert "slow_tests" in parsed
        assert "flaky_tests" in parsed
        assert "coverage_trend" in parsed
        assert "coverage_history" in parsed

    def test_format_empty_data(
        self,
        formatter: JSONFormatter,
    ) -> None:
        """Test JSON formatting with empty data."""
        data = {}
        result = formatter.format(data)

        parsed = json.loads(result)
        assert parsed == {}

    def test_format_indented_output(
        self,
        formatter: JSONFormatter,
    ) -> None:
        """Test that JSON output is properly indented."""
        data = {"coverage_trend": "improving"}
        result = formatter.format(data)

        # Should have newlines and indentation
        assert "\n" in result
        assert "  " in result


class TestHTMLFormatter:
    """Tests for HTMLFormatter class."""

    @pytest.fixture
    def formatter(self) -> HTMLFormatter:
        """Create HTMLFormatter instance."""
        return HTMLFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/test.py",
                current_coverage=60.0,
                target_coverage=90.0,
                lines_needed=30,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_slow",
                avg_duration=3.5,
                max_duration=7.0,
                min_duration=2.0,
                failure_rate=0.15,
                flakiness_score=0.6,
            ),
        ]

    def test_format_coverage_gaps_html(
        self,
        formatter: HTMLFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test HTML formatting of coverage gaps."""
        result = formatter._format_coverage_gaps_html(sample_gaps)

        assert "<h2>Coverage Gaps</h2>" in result
        assert "<table>" in result
        assert "src/test.py" in result
        assert "60.0%" in result
        assert "90.0%" in result
        assert "30" in result
        assert 'class="high"' in result
        assert "HIGH" in result

    def test_format_coverage_gaps_html_truncates_at_20(
        self,
        formatter: HTMLFormatter,
    ) -> None:
        """Test HTML coverage gaps are truncated to 20 items."""
        gaps = [
            CoverageGap(
                module=f"module_{i}.py",
                current_coverage=50.0,
                target_coverage=90.0,
                lines_needed=40,
                priority="high",
            )
            for i in range(30)
        ]

        result = formatter._format_coverage_gaps_html(gaps)

        assert "module_0.py" in result
        assert "module_19.py" in result
        assert "module_20.py" not in result

    def test_format_slow_tests_html(
        self,
        formatter: HTMLFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test HTML formatting of slow tests."""
        result = formatter._format_slow_tests_html(sample_tests)

        assert "<h2>Slowest Tests</h2>" in result
        assert "<table>" in result
        assert "test_slow" in result
        assert "3.50s" in result
        assert "7.00s" in result
        assert "15.0%" in result

    def test_format_flaky_tests_html(
        self,
        formatter: HTMLFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test HTML formatting of flaky tests."""
        result = formatter._format_flaky_tests_html(sample_tests)

        assert "<h2>Flaky Tests</h2>" in result
        assert "<table>" in result
        assert "test_slow" in result
        assert "0.600" in result
        assert "15.0%" in result

    def test_format_trends_html_improving(
        self,
        formatter: HTMLFormatter,
    ) -> None:
        """Test HTML formatting of improving trend."""
        history = [
            {"coverage": 85.0, "commit": "abc"},
            {"coverage": 80.0, "commit": "def"},
        ]

        result = formatter._format_trends_html("improving", history)

        assert "<h2>Coverage Trends</h2>" in result
        assert "📈" in result
        assert "IMPROVING" in result
        assert "85.0%" in result
        assert "+5.0%" in result

    def test_format_trends_html_declining(
        self,
        formatter: HTMLFormatter,
    ) -> None:
        """Test HTML formatting of declining trend."""
        history = [
            {"coverage": 75.0, "commit": "abc"},
            {"coverage": 80.0, "commit": "def"},
        ]

        result = formatter._format_trends_html("declining", history)

        assert "📉" in result
        assert "DECLINING" in result
        assert "-5.0%" in result

    def test_format_trends_html_stable(
        self,
        formatter: HTMLFormatter,
    ) -> None:
        """Test HTML formatting of stable trend."""
        history = [{"coverage": 85.0, "commit": "abc"}]

        result = formatter._format_trends_html("stable", history)

        assert "➡️" in result
        assert "STABLE" in result

    def test_format_trends_html_empty_history(
        self,
        formatter: HTMLFormatter,
    ) -> None:
        """Test HTML formatting with empty history."""
        result = formatter._format_trends_html("insufficient_data", [])

        assert "❓" in result
        assert "0.0%" in result

    def test_format_complete_html(
        self,
        formatter: HTMLFormatter,
        sample_gaps: list[CoverageGap],
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test complete HTML report generation."""
        data = {
            "coverage_gaps": sample_gaps,
            "slow_tests": sample_tests,
            "flaky_tests": sample_tests,
            "coverage_trend": "improving",
            "coverage_history": [
                {"coverage": 85.0, "commit": "abc"},
                {"coverage": 80.0, "commit": "def"},
            ],
        }

        result = formatter.format(data)

        # Check HTML structure
        assert "<!DOCTYPE html>" in result
        assert "<html>" in result
        assert "<head>" in result
        assert "<title>Test Analytics Report</title>" in result
        assert "<style>" in result
        assert "<body>" in result
        assert "<h1>Test Analytics Report</h1>" in result

        # Check all sections are included
        assert "Coverage Gaps" in result
        assert "Slowest Tests" in result
        assert "Flaky Tests" in result
        assert "Coverage Trends" in result

        # Check CSS is present
        assert "border-collapse" in result
        assert ".high" in result
        assert ".medium" in result
        assert ".low" in result

    def test_format_partial_html(
        self,
        formatter: HTMLFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test HTML formatting with only some data fields."""
        data = {"coverage_gaps": sample_gaps}

        result = formatter.format(data)

        assert "Coverage Gaps" in result
        assert "Slowest Tests" not in result
        assert "<!DOCTYPE html>" in result

    def test_format_empty_html(
        self,
        formatter: HTMLFormatter,
    ) -> None:
        """Test HTML formatting with no data."""
        data = {}
        result = formatter.format(data)

        # Should still have valid HTML structure
        assert "<!DOCTYPE html>" in result
        assert "<h1>Test Analytics Report</h1>" in result
        assert "</html>" in result


class TestYAMLFormatter:
    """Tests for YAMLFormatter class."""

    @pytest.fixture
    def formatter(self) -> YAMLFormatter:
        """Create YAMLFormatter instance."""
        return YAMLFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: YAMLFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test YAML formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        # Parse YAML and verify structure
        import yaml

        parsed = yaml.safe_load(result)

        assert "coverage_gaps" in parsed
        assert len(parsed["coverage_gaps"]) == 1
        assert parsed["coverage_gaps"][0]["module"] == "src/api/client.py"
        assert parsed["coverage_gaps"][0]["current_coverage"] == 85.5
        assert parsed["coverage_gaps"][0]["priority"] == "high"

    def test_format_slow_tests(
        self,
        formatter: YAMLFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test YAML formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        import yaml

        parsed = yaml.safe_load(result)

        assert "slow_tests" in parsed
        assert len(parsed["slow_tests"]) == 1
        assert parsed["slow_tests"][0]["test_name"] == "test_api_fetch"
        assert parsed["slow_tests"][0]["avg_duration"] == 2.45

    def test_format_empty_data(
        self,
        formatter: YAMLFormatter,
    ) -> None:
        """Test YAML formatting with empty data."""
        data = {}
        result = formatter.format(data)

        assert result == "{}\n"


class TestXMLFormatter:
    """Tests for XMLFormatter class."""

    @pytest.fixture
    def formatter(self) -> XMLFormatter:
        """Create XMLFormatter instance."""
        return XMLFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: XMLFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test XML formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        # Parse XML and verify structure
        root = ET.fromstring(
            result
        )  # noqa: S314  # parsing our own generated XML, not untrusted data

        assert root.tag == "test_analytics"
        gaps = root.find("coverage_gaps")
        assert gaps is not None
        assert len(gaps.findall("gap")) == 1

        gap = gaps.find("gap")
        assert gap is not None
        assert gap.find("module").text == "src/api/client.py"  # type: ignore[union-attr]
        assert gap.find("current_coverage").text == "85.5"  # type: ignore[union-attr]
        assert gap.find("priority").text == "high"  # type: ignore[union-attr]

    def test_format_slow_tests(
        self,
        formatter: XMLFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test XML formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        root = ET.fromstring(
            result
        )  # noqa: S314  # parsing our own generated XML, not untrusted data
        tests = root.find("slow_tests")
        assert tests is not None
        assert len(tests.findall("test")) == 1

    def test_format_empty_data(
        self,
        formatter: XMLFormatter,
    ) -> None:
        """Test XML formatting with empty data."""
        data = {}
        result = formatter.format(data)

        root = ET.fromstring(
            result
        )  # noqa: S314  # parsing our own generated XML, not untrusted data
        assert root.tag == "test_analytics"
        # Should have no child elements
        assert len(list(root)) == 0


class TestTableFormatter:
    """Tests for TableFormatter class."""

    @pytest.fixture
    def formatter(self) -> TableFormatter:
        """Create TableFormatter instance."""
        return TableFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: TableFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test table formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        assert "Coverage Gaps" in result
        assert "src/api/client.py" in result
        assert "85.5%" in result
        assert "HIGH" in result
        assert "+" in result  # Grid table borders

    def test_format_slow_tests(
        self,
        formatter: TableFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test table formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        assert "Slowest Tests" in result
        assert "test_api_fetch" in result
        assert "2.45s" in result

    def test_format_empty_data(
        self,
        formatter: TableFormatter,
    ) -> None:
        """Test table formatting with empty data."""
        data = {}
        result = formatter.format(data)

        # Empty string expected
        assert result == ""


class TestMarkdownFormatter:
    """Tests for MarkdownFormatter class."""

    @pytest.fixture
    def formatter(self) -> MarkdownFormatter:
        """Create MarkdownFormatter instance."""
        return MarkdownFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: MarkdownFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test markdown formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        assert "# Test Analytics Report" in result
        assert "## Coverage Gaps" in result
        assert "| Module |" in result
        assert "src/api/client.py" in result
        assert "85.5%" in result
        assert "HIGH" in result

    def test_format_slow_tests(
        self,
        formatter: MarkdownFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test markdown formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        assert "## Slowest Tests" in result
        assert "test_api_fetch" in result
        assert "2.45s" in result

    def test_format_trends(
        self,
        formatter: MarkdownFormatter,
    ) -> None:
        """Test markdown formatting of coverage trends."""
        data = {
            "coverage_trend": "improving",
            "coverage_history": [
                {"coverage": 90.0, "timestamp": "2026-05-01"},
                {"coverage": 85.0, "timestamp": "2026-04-01"},
            ],
        }
        result = formatter.format(data)

        assert "## Coverage Trends" in result
        assert "📈" in result
        assert "IMPROVING" in result
        assert "90.0%" in result

    def test_format_empty_data(
        self,
        formatter: MarkdownFormatter,
    ) -> None:
        """Test markdown formatting with empty data."""
        data = {}
        result = formatter.format(data)

        assert "# Test Analytics Report" in result


class TestCSVFormatter:
    """Tests for CSVFormatter class."""

    @pytest.fixture
    def formatter(self) -> CSVFormatter:
        """Create CSVFormatter instance."""
        return CSVFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: CSVFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test CSV formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        lines = result.strip().split("\n")

        # Check header
        assert "section" in lines[0]
        assert "module" in lines[0]

        # Check data
        assert "coverage_gaps" in lines[1]
        assert "src/api/client.py" in lines[1]
        assert "85.5" in lines[1]

    def test_format_slow_tests(
        self,
        formatter: CSVFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test CSV formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        lines = result.strip().split("\n")

        # Check data
        assert "slow_tests" in lines[1]
        assert "test_api_fetch" in lines[1]
        assert "2.45" in lines[1]

    def test_format_empty_data(
        self,
        formatter: CSVFormatter,
    ) -> None:
        """Test CSV formatting with empty data."""
        data = {}
        result = formatter.format(data)

        lines = result.strip().split("\n")
        # Should only have header
        assert len(lines) == 1
        assert "section" in lines[0]


class TestExcelFormatter:
    """Tests for ExcelFormatter class."""

    @pytest.fixture
    def formatter(self) -> ExcelFormatter:
        """Create ExcelFormatter instance."""
        return ExcelFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: ExcelFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test Excel formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        # Verify binary output
        assert isinstance(result, bytes)
        assert len(result) > 0

        # Parse Excel and verify structure
        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(result))
        assert "Coverage Gaps" in wb.sheetnames

        ws = wb["Coverage Gaps"]
        assert ws["A1"].value == "Module"
        assert ws["A2"].value == "src/api/client.py"
        assert ws["B2"].value == 85.5

    def test_format_slow_tests(
        self,
        formatter: ExcelFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test Excel formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(result))
        assert "Slow Tests" in wb.sheetnames

        ws = wb["Slow Tests"]
        assert ws["A1"].value == "Test"
        assert ws["A2"].value == "test_api_fetch"

    def test_format_multiple_sheets(
        self,
        formatter: ExcelFormatter,
        sample_gaps: list[CoverageGap],
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test Excel formatting with multiple sheets."""
        data = {
            "coverage_gaps": sample_gaps,
            "slow_tests": sample_tests,
            "flaky_tests": sample_tests,
        }
        result = formatter.format(data)

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(result))
        assert "Coverage Gaps" in wb.sheetnames
        assert "Slow Tests" in wb.sheetnames
        assert "Flaky Tests" in wb.sheetnames

    def test_format_empty_data(
        self,
        formatter: ExcelFormatter,
    ) -> None:
        """Test Excel formatting with empty data."""
        data = {}
        result = formatter.format(data)

        # Should still produce a valid Excel file
        assert isinstance(result, bytes)
        assert len(result) > 0


class TestTemplateFormatter:
    """Tests for TemplateFormatter class."""

    def test_format_with_template(
        self,
        tmp_path: Path,
    ) -> None:
        """Test template formatting with custom template."""
        template_file = tmp_path / "template.j2"
        template_file.write_text(
            "Gaps: {{ coverage_gaps|length }}\n" "Trend: {{ coverage_trend }}",
        )

        sample_gaps = [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

        formatter = TemplateFormatter(str(template_file))
        data = {"coverage_gaps": sample_gaps, "coverage_trend": "improving"}
        result = formatter.format(data)

        assert "Gaps: 1" in result
        assert "Trend: improving" in result

    def test_format_without_template(
        self,
    ) -> None:
        """Test template formatter requires template path."""
        formatter = TemplateFormatter()
        data = {"coverage_gaps": []}

        with pytest.raises(ValueError, match="Template path required"):
            formatter.format(data)

    def test_format_with_timestamp(
        self,
        tmp_path: Path,
    ) -> None:
        """Test template includes timestamp."""
        template_file = tmp_path / "template.j2"
        template_file.write_text("Time: {{ timestamp }}")

        formatter = TemplateFormatter(str(template_file))
        result = formatter.format({})

        assert "Time: " in result
        # Should have ISO format timestamp
        assert "T" in result  # ISO format separator

    def test_format_with_all_data(
        self,
        tmp_path: Path,
    ) -> None:
        """Test template with all data types."""
        template_file = tmp_path / "template.j2"
        template_file.write_text(
            "{% if coverage_gaps %}Gaps{% endif %}\n"
            "{% if slow_tests %}Slow{% endif %}\n"
            "{% if flaky_tests %}Flaky{% endif %}\n"
            "{% if coverage_trend %}Trend{% endif %}",
        )

        sample_gaps = [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

        sample_tests = [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

        formatter = TemplateFormatter(str(template_file))
        data = {
            "coverage_gaps": sample_gaps,
            "slow_tests": sample_tests,
            "flaky_tests": sample_tests,
            "coverage_trend": "improving",
        }
        result = formatter.format(data)

        assert "Gaps" in result
        assert "Slow" in result
        assert "Flaky" in result
        assert "Trend" in result
