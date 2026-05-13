"""Unit tests for ExcelFormatter class."""

from io import BytesIO

import pytest

from nhl_scrabble.analytics.analyzer import CoverageGap, TestPerformance
from nhl_scrabble.analytics.formatters import ExcelFormatter


class TestExcelFormatter:
    """Tests for ExcelFormatter class."""

    @pytest.fixture
    def formatter(self) -> ExcelFormatter:
        """Create ExcelFormatter instance."""
        return ExcelFormatter()

    @pytest.fixture
    def sample_gaps(self) -> list[CoverageGap]:
        """Create sample coverage gaps."""
        return [
            CoverageGap(
                module="src/api/client.py",
                current_coverage=85.5,
                target_coverage=90.0,
                lines_needed=12,
                priority="high",
            ),
        ]

    @pytest.fixture
    def sample_tests(self) -> list[TestPerformance]:
        """Create sample test performance data."""
        return [
            TestPerformance(
                test_name="test_api_fetch",
                avg_duration=2.45,
                max_duration=3.21,
                min_duration=1.89,
                failure_rate=0.05,
                flakiness_score=0.123,
            ),
        ]

    def test_format_coverage_gaps(
        self,
        formatter: ExcelFormatter,
        sample_gaps: list[CoverageGap],
    ) -> None:
        """Test Excel formatting of coverage gaps."""
        data = {"coverage_gaps": sample_gaps}
        result = formatter.format(data)

        # Verify binary output
        assert isinstance(result, bytes)
        assert len(result) > 0

        # Parse Excel and verify structure
        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(result))
        assert "Coverage Gaps" in wb.sheetnames

        ws = wb["Coverage Gaps"]
        assert ws["A1"].value == "Module"
        assert ws["A2"].value == "src/api/client.py"
        assert ws["B2"].value == 85.5

    def test_format_slow_tests(
        self,
        formatter: ExcelFormatter,
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test Excel formatting of slow tests."""
        data = {"slow_tests": sample_tests}
        result = formatter.format(data)

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(result))
        assert "Slow Tests" in wb.sheetnames

        ws = wb["Slow Tests"]
        assert ws["A1"].value == "Test"
        assert ws["A2"].value == "test_api_fetch"

    def test_format_multiple_sheets(
        self,
        formatter: ExcelFormatter,
        sample_gaps: list[CoverageGap],
        sample_tests: list[TestPerformance],
    ) -> None:
        """Test Excel formatting with multiple sheets."""
        data = {
            "coverage_gaps": sample_gaps,
            "slow_tests": sample_tests,
            "flaky_tests": sample_tests,
        }
        result = formatter.format(data)

        import openpyxl

        wb = openpyxl.load_workbook(BytesIO(result))
        assert "Coverage Gaps" in wb.sheetnames
        assert "Slow Tests" in wb.sheetnames
        assert "Flaky Tests" in wb.sheetnames

    def test_format_empty_data(
        self,
        formatter: ExcelFormatter,
    ) -> None:
        """Test Excel formatting with empty data."""
        data = {}
        result = formatter.format(data)

        # Should still produce a valid Excel file
        assert isinstance(result, bytes)
        assert len(result) > 0
