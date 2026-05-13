"""Excel formatter for test analytics data."""

from io import BytesIO
from typing import Any


class ExcelFormatter:
    """Format analytics data as Excel workbook.

    Creates a multi-sheet Excel workbook with separate sheets for coverage gaps,
    slow tests, flaky tests, and coverage trends.

    Note:
        Returns binary data (bytes), must be written in binary mode.

    Example:
        >>> formatter = ExcelFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> with open("analytics.xlsx", "wb") as f:
        ...     f.write(output)
    """

    def format(self, data: dict[str, Any]) -> bytes:  # type: ignore[explicit-any]  # noqa: C901, PLR0912  # Excel formatting requires multiple conditional branches
        """Format analytics data as Excel workbook.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            bytes: Excel workbook binary data.
        """
        import openpyxl  # noqa: PLC0415  # lazy import for optional dependency
        from openpyxl.styles import (  # noqa: PLC0415  # lazy import for optional dependency
            Font,
            PatternFill,
        )

        wb = openpyxl.Workbook()

        # Coverage Gaps sheet
        if data.get("coverage_gaps"):
            ws = wb.active
            if ws is not None:  # type: ignore[union-attr]  # openpyxl returns None when no sheets
                ws.title = "Coverage Gaps"
                ws.append(["Module", "Current", "Target", "Lines Needed", "Priority"])

                # Style header
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(
                        start_color="4CAF50",
                        end_color="4CAF50",
                        fill_type="solid",
                    )

                # Add data
                for gap in data["coverage_gaps"]:
                    ws.append(
                        [
                            gap.module,
                            gap.current_coverage,
                            gap.target_coverage,
                            gap.lines_needed,
                            gap.priority.upper(),
                        ],
                    )

        # Slow Tests sheet
        if data.get("slow_tests"):
            ws = wb.create_sheet("Slow Tests")
            ws.append(["Test", "Avg Duration", "Max Duration", "Failure Rate"])

            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="FFC107",
                    end_color="FFC107",
                    fill_type="solid",
                )

            # Add data
            for test in data["slow_tests"]:
                ws.append(
                    [
                        test.test_name,
                        test.avg_duration,
                        test.max_duration,
                        test.failure_rate,
                    ],
                )

        # Flaky Tests sheet
        if data.get("flaky_tests"):
            ws = wb.create_sheet("Flaky Tests")
            ws.append(["Test", "Flakiness Score", "Failure Rate"])

            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="F44336",
                    end_color="F44336",
                    fill_type="solid",
                )

            # Add data
            for test in data["flaky_tests"]:
                ws.append(
                    [
                        test.test_name,
                        test.flakiness_score,
                        test.failure_rate,
                    ],
                )

        # Coverage Trends sheet
        if data.get("coverage_history"):
            ws = wb.create_sheet("Coverage Trends")
            ws.append(["Timestamp", "Coverage"])

            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="2196F3",
                    end_color="2196F3",
                    fill_type="solid",
                )

            # Add data
            for record in data["coverage_history"]:
                ws.append(
                    [
                        record.get("timestamp", ""),
                        record.get("coverage", 0),
                    ],
                )

        # Save to bytes
        output = BytesIO()
        wb.save(output)
        return output.getvalue()
