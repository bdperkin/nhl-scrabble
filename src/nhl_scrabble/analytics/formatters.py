"""Output formatters for test analytics data."""

import csv
import json
import os
import xml.etree.ElementTree as ET
from datetime import UTC, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from rich.console import Console
from rich.panel import Panel
from rich.table import Table

from nhl_scrabble.analytics.analyzer import CoverageGap, TestPerformance


class TextFormatter:
    """Format analytics data as rich text output.

    Uses Rich library to create formatted tables and panels for terminal display.

    Example:
        >>> formatter = TextFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> print(output)
    """

    def __init__(self) -> None:
        """Initialize text formatter with Rich console."""
        self.console = Console()

    def _format_coverage_gaps(self, gaps: list[CoverageGap]) -> str:
        """Format coverage gaps as a table.

        Args:
            gaps: List of CoverageGap instances.

        Returns:
            str: Formatted table string.
        """
        table = Table(title="Coverage Gaps")
        table.add_column("Module", style="cyan")
        table.add_column("Current", justify="right")
        table.add_column("Target", justify="right")
        table.add_column("Lines Needed", justify="right")
        table.add_column("Priority", style="bold")

        for gap in gaps[:20]:  # Show top 20
            priority_color = {
                "high": "red",
                "medium": "yellow",
                "low": "green",
            }.get(gap.priority, "white")

            table.add_row(
                gap.module,
                f"{gap.current_coverage:.1f}%",
                f"{gap.target_coverage:.1f}%",
                str(gap.lines_needed),
                f"[{priority_color}]{gap.priority.upper()}[/{priority_color}]",
            )

        # Use console to render to string
        string_io = StringIO()
        temp_console = Console(file=string_io, force_terminal=True)
        temp_console.print(table)
        return string_io.getvalue()

    def _format_slow_tests(self, tests: list[TestPerformance]) -> str:
        """Format slow tests as a table.

        Args:
            tests: List of TestPerformance instances.

        Returns:
            str: Formatted table string.
        """
        table = Table(title="Slowest Tests")
        table.add_column("Test", style="cyan")
        table.add_column("Avg Duration", justify="right")
        table.add_column("Max Duration", justify="right")
        table.add_column("Failure Rate", justify="right")

        for test in tests:
            table.add_row(
                test.test_name,
                f"{test.avg_duration:.2f}s",
                f"{test.max_duration:.2f}s",
                f"{test.failure_rate * 100:.1f}%",
            )

        string_io = StringIO()
        temp_console = Console(file=string_io, force_terminal=True)
        temp_console.print(table)
        return string_io.getvalue()

    def _format_flaky_tests(self, tests: list[TestPerformance]) -> str:
        """Format flaky tests as a table.

        Args:
            tests: List of TestPerformance instances.

        Returns:
            str: Formatted table string.
        """
        table = Table(title="Flaky Tests")
        table.add_column("Test", style="cyan")
        table.add_column("Flakiness Score", justify="right")
        table.add_column("Failure Rate", justify="right")

        for test in tests:
            table.add_row(
                test.test_name,
                f"{test.flakiness_score:.3f}",
                f"{test.failure_rate * 100:.1f}%",
            )

        string_io = StringIO()
        temp_console = Console(file=string_io, force_terminal=True)
        temp_console.print(table)
        return string_io.getvalue()

    def _format_trends(  # type: ignore[explicit-any]
        self,
        trend: str,
        history: list[dict[str, Any]],
    ) -> str:
        """Format coverage trend information.

        Args:
            trend: Trend status ("improving", "declining", "stable", etc.).
            history: List of historical coverage data.

        Returns:
            str: Formatted trend panel string.
        """
        trend_emoji = {
            "improving": "📈",
            "declining": "📉",
            "stable": "➡️",
            "insufficient_data": "❓",
        }.get(trend, "")

        current = history[0]["coverage"] if history else 0
        oldest = history[-1]["coverage"] if len(history) > 1 else 0
        change = current - oldest

        text = f"{trend_emoji} Coverage Trend: {trend.upper()}\n"
        text += f"Current: {current:.1f}%\n"
        text += f"Change (30 days): {change:+.1f}%"

        panel = Panel(text, title="Coverage Trends", border_style="blue")

        string_io = StringIO()
        temp_console = Console(file=string_io, force_terminal=True)
        temp_console.print(panel)
        return string_io.getvalue()

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as rich text.

        Args:
            data: Dictionary containing analytics data with keys like
                "coverage_gaps", "slow_tests", "flaky_tests", "coverage_trend".

        Returns:
            str: Formatted text output ready for display.
        """
        output = []

        # Coverage gaps
        if "coverage_gaps" in data:
            output.append(self._format_coverage_gaps(data["coverage_gaps"]))

        # Slow tests
        if "slow_tests" in data:
            output.append(self._format_slow_tests(data["slow_tests"]))

        # Flaky tests
        if "flaky_tests" in data:
            output.append(self._format_flaky_tests(data["flaky_tests"]))

        # Trends
        if "coverage_trend" in data:
            output.append(
                self._format_trends(
                    data["coverage_trend"],
                    data.get("coverage_history", []),
                ),
            )

        return "\n\n".join(output)


class JSONFormatter:
    """Format analytics data as JSON.

    Converts analytics data structures to JSON for programmatic consumption.

    Example:
        >>> formatter = JSONFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> data = json.loads(output)
    """

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as JSON.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: JSON-formatted string.
        """
        # Convert dataclasses to dicts for JSON serialization
        serializable_data = {}

        if "coverage_gaps" in data:
            serializable_data["coverage_gaps"] = [
                {
                    "module": gap.module,
                    "current_coverage": gap.current_coverage,
                    "target_coverage": gap.target_coverage,
                    "lines_needed": gap.lines_needed,
                    "priority": gap.priority,
                }
                for gap in data["coverage_gaps"]
            ]

        if "slow_tests" in data:
            serializable_data["slow_tests"] = [
                {
                    "test_name": test.test_name,
                    "avg_duration": test.avg_duration,
                    "max_duration": test.max_duration,
                    "min_duration": test.min_duration,
                    "failure_rate": test.failure_rate,
                    "flakiness_score": test.flakiness_score,
                }
                for test in data["slow_tests"]
            ]

        if "flaky_tests" in data:
            serializable_data["flaky_tests"] = [
                {
                    "test_name": test.test_name,
                    "flakiness_score": test.flakiness_score,
                    "failure_rate": test.failure_rate,
                    "avg_duration": test.avg_duration,
                }
                for test in data["flaky_tests"]
            ]

        if "coverage_trend" in data:
            serializable_data["coverage_trend"] = data["coverage_trend"]

        if "coverage_history" in data:
            serializable_data["coverage_history"] = data["coverage_history"]

        return json.dumps(serializable_data, indent=2, default=str)


class HTMLFormatter:
    """Format analytics data as HTML.

    Generates HTML tables and styled output for web display.

    Example:
        >>> formatter = HTMLFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> Path("report.html").write_text(output)
    """

    def _format_coverage_gaps_html(self, gaps: list[CoverageGap]) -> str:
        """Format coverage gaps as HTML table.

        Args:
            gaps: List of CoverageGap instances.

        Returns:
            str: HTML table string.
        """
        html = ["<h2>Coverage Gaps</h2>", "<table>", "<tr>"]
        html.append("<th>Module</th>")
        html.append("<th>Current</th>")
        html.append("<th>Target</th>")
        html.append("<th>Lines Needed</th>")
        html.append("<th>Priority</th>")
        html.append("</tr>")

        for gap in gaps[:20]:  # Show top 20
            html.append("<tr>")
            html.append(f"<td>{gap.module}</td>")
            html.append(f"<td>{gap.current_coverage:.1f}%</td>")
            html.append(f"<td>{gap.target_coverage:.1f}%</td>")
            html.append(f"<td>{gap.lines_needed}</td>")
            html.append(
                f'<td class="{gap.priority}">{gap.priority.upper()}</td>',
            )
            html.append("</tr>")

        html.append("</table>")
        return "\n".join(html)

    def _format_slow_tests_html(self, tests: list[TestPerformance]) -> str:
        """Format slow tests as HTML table.

        Args:
            tests: List of TestPerformance instances.

        Returns:
            str: HTML table string.
        """
        html = ["<h2>Slowest Tests</h2>", "<table>", "<tr>"]
        html.append("<th>Test</th>")
        html.append("<th>Avg Duration</th>")
        html.append("<th>Max Duration</th>")
        html.append("<th>Failure Rate</th>")
        html.append("</tr>")

        for test in tests:
            html.append("<tr>")
            html.append(f"<td>{test.test_name}</td>")
            html.append(f"<td>{test.avg_duration:.2f}s</td>")
            html.append(f"<td>{test.max_duration:.2f}s</td>")
            html.append(f"<td>{test.failure_rate * 100:.1f}%</td>")
            html.append("</tr>")

        html.append("</table>")
        return "\n".join(html)

    def _format_flaky_tests_html(self, tests: list[TestPerformance]) -> str:
        """Format flaky tests as HTML table.

        Args:
            tests: List of TestPerformance instances.

        Returns:
            str: HTML table string.
        """
        html = ["<h2>Flaky Tests</h2>", "<table>", "<tr>"]
        html.append("<th>Test</th>")
        html.append("<th>Flakiness Score</th>")
        html.append("<th>Failure Rate</th>")
        html.append("</tr>")

        for test in tests:
            html.append("<tr>")
            html.append(f"<td>{test.test_name}</td>")
            html.append(f"<td>{test.flakiness_score:.3f}</td>")
            html.append(f"<td>{test.failure_rate * 100:.1f}%</td>")
            html.append("</tr>")

        html.append("</table>")
        return "\n".join(html)

    def _format_trends_html(  # type: ignore[explicit-any]
        self,
        trend: str,
        history: list[dict[str, Any]],
    ) -> str:
        """Format coverage trend as HTML.

        Args:
            trend: Trend status.
            history: Historical coverage data.

        Returns:
            str: HTML div string.
        """
        trend_emoji = {
            "improving": "📈",
            "declining": "📉",
            "stable": "➡️",
            "insufficient_data": "❓",
        }.get(trend, "")

        current = history[0]["coverage"] if history else 0
        oldest = history[-1]["coverage"] if len(history) > 1 else 0
        change = current - oldest

        html = ['<div class="trend">', "<h2>Coverage Trends</h2>"]
        html.append(f"<p>{trend_emoji} <strong>Trend:</strong> {trend.upper()}</p>")
        html.append(f"<p><strong>Current:</strong> {current:.1f}%</p>")
        html.append(f"<p><strong>Change (30 days):</strong> {change:+.1f}%</p>")
        html.append("</div>")

        return "\n".join(html)

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as HTML.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: HTML-formatted string.
        """
        html_parts = [
            "<!DOCTYPE html>",
            "<html>",
            "<head>",
            "<title>Test Analytics Report</title>",
            "<style>",
            "body { font-family: Arial, sans-serif; margin: 20px; }",
            "h1 { color: #333; }",
            "h2 { color: #666; margin-top: 30px; }",
            "table { border-collapse: collapse; width: 100%; margin: 20px 0; }",
            "th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }",
            "th { background-color: #4CAF50; color: white; }",
            "tr:nth-child(even) { background-color: #f2f2f2; }",
            ".high { color: red; font-weight: bold; }",
            ".medium { color: orange; font-weight: bold; }",
            ".low { color: green; font-weight: bold; }",
            ".trend { padding: 20px; background-color: #e7f3ff; border-left: 4px solid #2196F3; margin: 20px 0; }",
            "</style>",
            "</head>",
            "<body>",
            "<h1>Test Analytics Report</h1>",
        ]

        # Coverage gaps
        if "coverage_gaps" in data:
            html_parts.append(self._format_coverage_gaps_html(data["coverage_gaps"]))

        # Slow tests
        if "slow_tests" in data:
            html_parts.append(self._format_slow_tests_html(data["slow_tests"]))

        # Flaky tests
        if "flaky_tests" in data:
            html_parts.append(self._format_flaky_tests_html(data["flaky_tests"]))

        # Trends
        if "coverage_trend" in data:
            html_parts.append(
                self._format_trends_html(
                    data["coverage_trend"],
                    data.get("coverage_history", []),
                ),
            )

        html_parts.extend(["</body>", "</html>"])

        return "\n".join(html_parts)


class YAMLFormatter:
    """Format analytics data as YAML.

    Outputs analytics data in YAML format for configuration files and
    human-readable structured data.

    Example:
        >>> formatter = YAMLFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> with open("analytics.yaml", "w") as f:
        ...     f.write(output)
    """

    def _convert_to_dict(self, data: dict[str, Any]) -> dict[str, Any]:  # type: ignore[explicit-any]
        """Convert dataclasses to dicts for serialization.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            dict: Serializable dictionary.
        """
        serializable_data: dict[str, Any] = {}  # type: ignore[explicit-any]

        if "coverage_gaps" in data:
            serializable_data["coverage_gaps"] = [
                {
                    "module": gap.module,
                    "current_coverage": gap.current_coverage,
                    "target_coverage": gap.target_coverage,
                    "lines_needed": gap.lines_needed,
                    "priority": gap.priority,
                }
                for gap in data["coverage_gaps"]
            ]

        if "slow_tests" in data:
            serializable_data["slow_tests"] = [
                {
                    "test_name": test.test_name,
                    "avg_duration": test.avg_duration,
                    "max_duration": test.max_duration,
                    "min_duration": test.min_duration,
                    "failure_rate": test.failure_rate,
                    "flakiness_score": test.flakiness_score,
                }
                for test in data["slow_tests"]
            ]

        if "flaky_tests" in data:
            serializable_data["flaky_tests"] = [
                {
                    "test_name": test.test_name,
                    "flakiness_score": test.flakiness_score,
                    "failure_rate": test.failure_rate,
                    "avg_duration": test.avg_duration,
                }
                for test in data["flaky_tests"]
            ]

        if "coverage_trend" in data:
            serializable_data["coverage_trend"] = data["coverage_trend"]

        if "coverage_history" in data:
            serializable_data["coverage_history"] = data["coverage_history"]

        return serializable_data

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as YAML.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: YAML-formatted string.
        """
        import yaml  # noqa: PLC0415  # lazy import for optional dependency

        serializable_data = self._convert_to_dict(data)
        return yaml.dump(serializable_data, default_flow_style=False, sort_keys=False)


class XMLFormatter:
    """Format analytics data as XML.

    Outputs analytics data in XML format for enterprise systems and
    legacy integrations.

    Example:
        >>> formatter = XMLFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> with open("analytics.xml", "w") as f:
        ...     f.write(output)
    """

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as XML.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: XML-formatted string.
        """
        root = ET.Element("test_analytics")

        # Add coverage gaps
        if "coverage_gaps" in data:
            gaps_elem = ET.SubElement(root, "coverage_gaps")
            for gap in data["coverage_gaps"]:
                gap_elem = ET.SubElement(gaps_elem, "gap")
                ET.SubElement(gap_elem, "module").text = gap.module
                ET.SubElement(gap_elem, "current_coverage").text = str(
                    gap.current_coverage,
                )
                ET.SubElement(gap_elem, "target_coverage").text = str(
                    gap.target_coverage,
                )
                ET.SubElement(gap_elem, "lines_needed").text = str(gap.lines_needed)
                ET.SubElement(gap_elem, "priority").text = gap.priority

        # Add slow tests
        if "slow_tests" in data:
            tests_elem = ET.SubElement(root, "slow_tests")
            for test in data["slow_tests"]:
                test_elem = ET.SubElement(tests_elem, "test")
                ET.SubElement(test_elem, "test_name").text = test.test_name
                ET.SubElement(test_elem, "avg_duration").text = str(test.avg_duration)
                ET.SubElement(test_elem, "max_duration").text = str(test.max_duration)
                ET.SubElement(test_elem, "min_duration").text = str(test.min_duration)
                ET.SubElement(test_elem, "failure_rate").text = str(test.failure_rate)
                ET.SubElement(test_elem, "flakiness_score").text = str(
                    test.flakiness_score,
                )

        # Add flaky tests
        if "flaky_tests" in data:
            flaky_elem = ET.SubElement(root, "flaky_tests")
            for test in data["flaky_tests"]:
                test_elem = ET.SubElement(flaky_elem, "test")
                ET.SubElement(test_elem, "test_name").text = test.test_name
                ET.SubElement(test_elem, "flakiness_score").text = str(
                    test.flakiness_score,
                )
                ET.SubElement(test_elem, "failure_rate").text = str(test.failure_rate)
                ET.SubElement(test_elem, "avg_duration").text = str(test.avg_duration)

        # Add coverage trend
        if "coverage_trend" in data:
            trend_elem = ET.SubElement(root, "coverage_trend")
            trend_elem.text = data["coverage_trend"]

        # Add coverage history
        if "coverage_history" in data:
            history_elem = ET.SubElement(root, "coverage_history")
            for record in data["coverage_history"]:
                record_elem = ET.SubElement(history_elem, "record")
                ET.SubElement(record_elem, "timestamp").text = str(
                    record.get("timestamp", ""),
                )
                ET.SubElement(record_elem, "coverage").text = str(
                    record.get("coverage", 0),
                )

        return ET.tostring(root, encoding="unicode", method="xml")


class TableFormatter:
    """Format analytics data as simple ASCII tables.

    Outputs analytics data in simple ASCII table format using the
    tabulate library (simpler than Rich text output).

    Example:
        >>> formatter = TableFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> print(output)
    """

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as ASCII tables.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: ASCII table-formatted string.
        """
        from tabulate import tabulate  # noqa: PLC0415  # lazy import for optional dependency

        output = []

        # Coverage gaps table
        if "coverage_gaps" in data:
            headers = ["Module", "Current", "Target", "Lines Needed", "Priority"]
            rows = [
                [
                    gap.module,
                    f"{gap.current_coverage:.1f}%",
                    f"{gap.target_coverage:.1f}%",
                    gap.lines_needed,
                    gap.priority.upper(),
                ]
                for gap in data["coverage_gaps"][:20]
            ]
            output.append("Coverage Gaps\n")
            output.append(tabulate(rows, headers=headers, tablefmt="grid"))

        # Slow tests table
        if "slow_tests" in data:
            headers = ["Test", "Avg Duration", "Max Duration", "Failure Rate"]
            rows = [
                [
                    test.test_name,
                    f"{test.avg_duration:.2f}s",
                    f"{test.max_duration:.2f}s",
                    f"{test.failure_rate * 100:.1f}%",
                ]
                for test in data["slow_tests"]
            ]
            output.append("Slowest Tests\n")
            output.append(tabulate(rows, headers=headers, tablefmt="grid"))

        # Flaky tests table
        if "flaky_tests" in data:
            headers = ["Test", "Flakiness Score", "Failure Rate"]
            rows = [
                [
                    test.test_name,
                    f"{test.flakiness_score:.3f}",
                    f"{test.failure_rate * 100:.1f}%",
                ]
                for test in data["flaky_tests"]
            ]
            output.append("Flaky Tests\n")
            output.append(tabulate(rows, headers=headers, tablefmt="grid"))

        return "\n\n".join(output)


class MarkdownFormatter:
    """Format analytics data as Markdown.

    Outputs analytics data in Markdown format for documentation and GitHub.

    Example:
        >>> formatter = MarkdownFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> with open("analytics.md", "w") as f:
        ...     f.write(output)
    """

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as Markdown.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: Markdown-formatted string.
        """
        output = ["# Test Analytics Report\n"]

        # Coverage gaps
        if "coverage_gaps" in data:
            output.append("## Coverage Gaps\n")
            output.append("| Module | Current | Target | Lines Needed | Priority |")
            output.append("|--------|---------|--------|--------------|----------|")
            output.extend(
                [
                    f"| {gap.module} | {gap.current_coverage:.1f}% | "
                    f"{gap.target_coverage:.1f}% | {gap.lines_needed} | "
                    f"{gap.priority.upper()} |"
                    for gap in data["coverage_gaps"][:20]
                ],
            )
            output.append("")

        # Slow tests
        if "slow_tests" in data:
            output.append("## Slowest Tests\n")
            output.append("| Test | Avg Duration | Max Duration | Failure Rate |")
            output.append("|------|--------------|--------------|--------------|")
            output.extend(
                [
                    f"| {test.test_name} | {test.avg_duration:.2f}s | "
                    f"{test.max_duration:.2f}s | {test.failure_rate * 100:.1f}% |"
                    for test in data["slow_tests"]
                ],
            )
            output.append("")

        # Flaky tests
        if "flaky_tests" in data:
            output.append("## Flaky Tests\n")
            output.append("| Test | Flakiness Score | Failure Rate |")
            output.append("|------|-----------------|--------------|")
            output.extend(
                [
                    f"| {test.test_name} | {test.flakiness_score:.3f} | "
                    f"{test.failure_rate * 100:.1f}% |"
                    for test in data["flaky_tests"]
                ],
            )
            output.append("")

        # Coverage trend
        if "coverage_trend" in data:
            output.append("## Coverage Trends\n")
            trend = data["coverage_trend"]
            history = data.get("coverage_history", [])

            trend_emoji = {
                "improving": "📈",
                "declining": "📉",
                "stable": "➡️",
                "insufficient_data": "❓",
            }.get(trend, "")

            output.append(f"{trend_emoji} **Trend:** {trend.upper()}\n")

            if history:
                current = history[0]["coverage"]
                oldest = history[-1]["coverage"] if len(history) > 1 else 0
                change = current - oldest
                output.append(f"- **Current:** {current:.1f}%")
                output.append(f"- **Change (30 days):** {change:+.1f}%\n")

        return "\n".join(output)


class CSVFormatter:
    """Format analytics data as CSV.

    Outputs analytics data in CSV format for spreadsheet import and data analysis.

    Example:
        >>> formatter = CSVFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> with open("analytics.csv", "w") as f:
        ...     f.write(output)
    """

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data as CSV.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: CSV-formatted string.
        """
        output = StringIO()
        writer = csv.writer(output)

        # Write header
        writer.writerow(
            [
                "section",
                "test_name",
                "module",
                "current_coverage",
                "target_coverage",
                "lines_needed",
                "priority",
                "avg_duration",
                "max_duration",
                "min_duration",
                "failure_rate",
                "flakiness_score",
            ],
        )

        # Write coverage gaps
        if "coverage_gaps" in data:
            for gap in data["coverage_gaps"]:
                writer.writerow(
                    [
                        "coverage_gaps",
                        "",
                        gap.module,
                        gap.current_coverage,
                        gap.target_coverage,
                        gap.lines_needed,
                        gap.priority,
                        "",
                        "",
                        "",
                        "",
                        "",
                    ],
                )

        # Write slow tests
        if "slow_tests" in data:
            for test in data["slow_tests"]:
                writer.writerow(
                    [
                        "slow_tests",
                        test.test_name,
                        "",
                        "",
                        "",
                        "",
                        "",
                        test.avg_duration,
                        test.max_duration,
                        test.min_duration,
                        test.failure_rate,
                        test.flakiness_score,
                    ],
                )

        # Write flaky tests
        if "flaky_tests" in data:
            for test in data["flaky_tests"]:
                writer.writerow(
                    [
                        "flaky_tests",
                        test.test_name,
                        "",
                        "",
                        "",
                        "",
                        "",
                        test.avg_duration,
                        "",
                        "",
                        test.failure_rate,
                        test.flakiness_score,
                    ],
                )

        return output.getvalue()


class ExcelFormatter:
    """Format analytics data as Excel workbook.

    Creates a multi-sheet Excel workbook with separate sheets for coverage gaps,
    slow tests, flaky tests, and coverage trends.

    Note:
        Returns binary data (bytes), must be written in binary mode.

    Example:
        >>> formatter = ExcelFormatter()
        >>> output = formatter.format({"coverage_gaps": gaps})
        >>> with open("analytics.xlsx", "wb") as f:
        ...     f.write(output)
    """

    def format(self, data: dict[str, Any]) -> bytes:  # type: ignore[explicit-any]  # noqa: C901, PLR0912  # Excel formatting requires multiple conditional branches
        """Format analytics data as Excel workbook.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            bytes: Excel workbook binary data.
        """
        import openpyxl  # noqa: PLC0415  # lazy import for optional dependency
        from openpyxl.styles import (  # noqa: PLC0415  # lazy import for optional dependency
            Font,
            PatternFill,
        )

        wb = openpyxl.Workbook()

        # Coverage Gaps sheet
        if data.get("coverage_gaps"):
            ws = wb.active
            if ws is not None:  # type: ignore[union-attr]  # openpyxl returns None when no sheets
                ws.title = "Coverage Gaps"
                ws.append(["Module", "Current", "Target", "Lines Needed", "Priority"])

                # Style header
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(
                        start_color="4CAF50",
                        end_color="4CAF50",
                        fill_type="solid",
                    )

                # Add data
                for gap in data["coverage_gaps"]:
                    ws.append(
                        [
                            gap.module,
                            gap.current_coverage,
                            gap.target_coverage,
                            gap.lines_needed,
                            gap.priority.upper(),
                        ],
                    )

        # Slow Tests sheet
        if data.get("slow_tests"):
            ws = wb.create_sheet("Slow Tests")
            ws.append(["Test", "Avg Duration", "Max Duration", "Failure Rate"])

            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="FFC107",
                    end_color="FFC107",
                    fill_type="solid",
                )

            # Add data
            for test in data["slow_tests"]:
                ws.append(
                    [
                        test.test_name,
                        test.avg_duration,
                        test.max_duration,
                        test.failure_rate,
                    ],
                )

        # Flaky Tests sheet
        if data.get("flaky_tests"):
            ws = wb.create_sheet("Flaky Tests")
            ws.append(["Test", "Flakiness Score", "Failure Rate"])

            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="F44336",
                    end_color="F44336",
                    fill_type="solid",
                )

            # Add data
            for test in data["flaky_tests"]:
                ws.append(
                    [
                        test.test_name,
                        test.flakiness_score,
                        test.failure_rate,
                    ],
                )

        # Coverage Trends sheet
        if data.get("coverage_history"):
            ws = wb.create_sheet("Coverage Trends")
            ws.append(["Timestamp", "Coverage"])

            # Style header
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="2196F3",
                    end_color="2196F3",
                    fill_type="solid",
                )

            # Add data
            for record in data["coverage_history"]:
                ws.append(
                    [
                        record.get("timestamp", ""),
                        record.get("coverage", 0),
                    ],
                )

        # Save to bytes
        output = BytesIO()
        wb.save(output)
        return output.getvalue()


class TemplateFormatter:
    """Format analytics data using custom Jinja2 template.

    Allows users to provide a custom Jinja2 template file for flexible
    output formatting.

    Template variables available:
        - coverage_gaps: List of coverage gap data
        - slow_tests: List of slow test data
        - flaky_tests: List of flaky test data
        - coverage_trend: Trend status string
        - coverage_history: Historical coverage data
        - timestamp: Current timestamp

    Example:
        >>> formatter = TemplateFormatter("/path/to/template.j2")
        >>> output = formatter.format({"coverage_gaps": gaps})
    """

    def __init__(self, template_path: str | None = None) -> None:
        """Initialize template formatter.

        Args:
            template_path: Path to Jinja2 template file. If not provided,
                will check NHL_SCRABBLE_ANALYTICS_TEMPLATE environment variable.
        """
        self.template_path = template_path or os.getenv(
            "NHL_SCRABBLE_ANALYTICS_TEMPLATE",
        )

    def format(self, data: dict[str, Any]) -> str:  # type: ignore[explicit-any]
        """Format analytics data using template.

        Args:
            data: Dictionary containing analytics data.

        Returns:
            str: Formatted output from template.

        Raises:
            ValueError: If template path is not provided.
        """
        from jinja2 import (  # noqa: PLC0415  # lazy import for optional dependency
            Environment,
            FileSystemLoader,
            select_autoescape,
        )

        if not self.template_path:
            msg = (
                "Template path required. Set NHL_SCRABBLE_ANALYTICS_TEMPLATE "
                "or use --template option."
            )
            raise ValueError(msg)

        template_path = Path(self.template_path)
        env = Environment(
            loader=FileSystemLoader(template_path.parent),
            autoescape=select_autoescape(),  # Enable autoescape for security
        )
        template = env.get_template(template_path.name)  # clearer without chaining

        # Convert dataclasses to dicts for template rendering

        render_data: dict[str, Any] = {"timestamp": datetime.now(tz=UTC).isoformat()}  # type: ignore[explicit-any]

        if "coverage_gaps" in data:
            render_data["coverage_gaps"] = [
                {
                    "module": gap.module,
                    "current_coverage": gap.current_coverage,
                    "target_coverage": gap.target_coverage,
                    "lines_needed": gap.lines_needed,
                    "priority": gap.priority,
                }
                for gap in data["coverage_gaps"]
            ]

        if "slow_tests" in data:
            render_data["slow_tests"] = [
                {
                    "test_name": test.test_name,
                    "avg_duration": test.avg_duration,
                    "max_duration": test.max_duration,
                    "min_duration": test.min_duration,
                    "failure_rate": test.failure_rate,
                    "flakiness_score": test.flakiness_score,
                }
                for test in data["slow_tests"]
            ]

        if "flaky_tests" in data:
            render_data["flaky_tests"] = [
                {
                    "test_name": test.test_name,
                    "flakiness_score": test.flakiness_score,
                    "failure_rate": test.failure_rate,
                    "avg_duration": test.avg_duration,
                }
                for test in data["flaky_tests"]
            ]

        if "coverage_trend" in data:
            render_data["coverage_trend"] = data["coverage_trend"]

        if "coverage_history" in data:
            render_data["coverage_history"] = data["coverage_history"]

        return template.render(**render_data)
